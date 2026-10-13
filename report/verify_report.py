"""Check the built report against the submission guidelines.

    cd backend && .venv/Scripts/python ../report/verify_report.py

Verifies the font/size/spacing rules, that every figure and table is referred to
in the text, that no reference points at a figure or table that does not exist,
and reports the page count measured by Word (report/pagemap.json).
"""

from __future__ import annotations

import io
import json
import os
import re
from collections import Counter

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
DOCX = os.path.join(HERE, "VibeShift_MiniProject_Report.docx")

problems: list[str] = []
notes: list[str] = []


def check(condition: bool, message: str) -> None:
    if not condition:
        problems.append(message)


def main() -> None:
    d = docx.Document(DOCX)

    # ---- page geometry --------------------------------------------------- #
    for section in d.sections:
        check(abs(section.left_margin.inches - 1) < 0.01, "left margin must be 1 inch")
        check(abs(section.right_margin.inches - 1) < 0.01, "right margin must be 1 inch")
        check(abs(section.top_margin.inches - 1) < 0.01, "top margin must be 1 inch")
        check(abs(section.bottom_margin.inches - 1) < 0.01, "bottom margin must be 1 inch")

    # ---- font, size, spacing, alignment ---------------------------------- #
    expected = {
        "Normal": (12, False, True),       # size, bold, justified
        "Heading 1": (16, True, False),
        "Heading 2": (14, True, False),
        "Heading 3": (12, True, False),
    }
    for name, (size, bold, justify) in expected.items():
        style = d.styles[name]
        check(style.font.name == "Times New Roman", f"{name} font must be Times New Roman")
        check(style.font.size == Pt(size), f"{name} must be {size} pt")
        check(bool(style.font.bold) == bold, f"{name} bold must be {bold}")
        if justify:
            check(style.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY,
                  "body text must be justified")
        check(abs(float(style.paragraph_format.line_spacing) - 1.5) < 0.001,
              f"{name} line spacing must be 1.5")

    body = [p for p in d.paragraphs if p.style.name == "Normal" and p.text.strip()]
    notes.append(f"body paragraphs: {len(body)}")

    # ---- captions and cross references ----------------------------------- #
    caption_re = re.compile(r"^(Figure|Table) (\d+\.\d+):\s*(.+)$")
    captions: dict[str, dict[str, str]] = {"Figure": {}, "Table": {}}
    for p in d.paragraphs:
        m = caption_re.match(p.text.strip())
        if m and p.style.name == "Caption":
            captions[m.group(1)][m.group(2)] = m.group(3)

    text = "\n".join(p.text for p in d.paragraphs)
    for kind, items in captions.items():
        for key in items:
            # the caption itself contains the number, so require a second mention
            mentions = len(re.findall(rf"{kind}\s+{re.escape(key)}\b", text))
            check(mentions >= 2, f"{kind} {key} is never referred to in the text")
        for ref in set(re.findall(rf"{kind}\s+(\d+\.\d+)", text)):
            check(ref in items, f"text refers to {kind} {ref}, which does not exist")

    notes.append(f"figures: {len(captions['Figure'])}  tables: {len(captions['Table'])}")

    # ---- numbering must be per chapter and gap free ---------------------- #
    for kind, items in captions.items():
        by_chapter: dict[int, list[int]] = {}
        for key in items:
            chapter, minor = (int(x) for x in key.split("."))
            by_chapter.setdefault(chapter, []).append(minor)
        for chapter, minors in sorted(by_chapter.items()):
            check(sorted(minors) == list(range(1, len(minors) + 1)),
                  f"{kind} numbering in chapter {chapter} is not 1..n: {sorted(minors)}")

    # ---- tables use the required fonts ----------------------------------- #
    table_fonts = Counter()
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        if run.font.name or run.font.size:
                            table_fonts[(run.font.name, run.font.size.pt if run.font.size else None)] += 1
    bad = [k for k in table_fonts if k[0] not in (None, "Times New Roman")]
    check(not bad, f"table runs with a non-Times font: {bad}")

    # ---- page count from the Word measurement ---------------------------- #
    pm_path = os.path.join(HERE, "pagemap.json")
    if os.path.exists(pm_path):
        with io.open(pm_path, encoding="utf-8-sig") as fh:
            pm = json.load(fh)
        notes.append(f"Word page count: {pm.get('pages')}")
        check(int(pm.get("pages") or 0) <= 19, "report must stay under 20 pages")
        missing = [k for k in captions["Figure"] if str(k) not in pm.get("figures", {})]
        check(not missing, f"figures without a measured page number: {missing}")
        missing = [k for k in captions["Table"] if str(k) not in pm.get("tables", {})]
        check(not missing, f"tables without a measured page number: {missing}")
    else:
        notes.append("pagemap.json missing - run report/word_pagemap.ps1 to measure pages")

    print("report verification")
    for note in notes:
        print(f"  - {note}")
    if problems:
        print("\nPROBLEMS")
        for p in problems:
            print(f"  x {p}")
        raise SystemExit(1)
    print("\nall guideline checks passed")


if __name__ == "__main__":
    main()
